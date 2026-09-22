"""
Checagem de integridade: existe alguma linha duplicada (mesmo Auto de
Infração aparecendo mais de uma vez) na planilha FINAL do SharePoint?

Por que existe (22/09/2026, a pedido do usuário - "precisamos ser
perfeitos"): a proteção contra duplicata (`planilha.ja_registrado()`) nunca
foi confirmada em escala real com múltiplos workers consolidando na mesma
planilha final (`scripts/consolidar_planilhas.py`) - isso é uma suposição
de design, não um fato verificado. Só lê a planilha final, sem sessão nem
contato com o portal - seguro rodar com qualquer varredura real em
andamento em paralelo.

USO:
    python scripts/checar_duplicidade_planilha.py
"""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR / "src"))

from robo_antt.config import PLANILHA_PATH  # noqa: E402
from robo_antt.planilha import COLUNAS, NOME_ABA  # noqa: E402

_COL_AUTO = COLUNAS.index("Auto de Infração")


def main() -> None:
    print(f"Planilha final: {PLANILHA_PATH}\n")

    if not PLANILHA_PATH.exists():
        print("[ATENCAO] Planilha final ainda nao existe - nada pra checar.")
        return

    wb = load_workbook(str(PLANILHA_PATH), read_only=True)
    ws = wb[NOME_ABA] if NOME_ABA in wb.sheetnames else wb.active

    linhas_por_auto: dict[str, list[int]] = defaultdict(list)
    total_linhas = 0
    linhas_sem_auto = 0
    for numero_linha, linha in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        total_linhas += 1
        auto = linha[_COL_AUTO] if len(linha) > _COL_AUTO else None
        if not auto:
            linhas_sem_auto += 1
            continue
        linhas_por_auto[str(auto)].append(numero_linha)
    wb.close()

    duplicados = {auto: nums for auto, nums in linhas_por_auto.items() if len(nums) > 1}

    print(f"Total de linhas de dados: {total_linhas}")
    print(f"Autos distintos: {len(linhas_por_auto)}")
    if linhas_sem_auto:
        print(f"[ATENCAO] {linhas_sem_auto} linha(s) sem valor em 'Auto de Infracao' (celula vazia).")

    if not duplicados:
        print("\n[OK] Nenhum Auto de Infracao duplicado - cada linha da planilha final e unica.")
        return

    total_linhas_extras = sum(len(nums) - 1 for nums in duplicados.values())
    print(
        f"\n[ATENCAO] {len(duplicados)} Auto(s) de Infracao aparecem em mais de 1 linha "
        f"({total_linhas_extras} linha(s) extra(s) no total):"
    )
    for auto, nums in sorted(duplicados.items()):
        print(f"  - {auto}: linhas {nums}")

    print(
        "\nCada duplicata acima precisa de remocao manual das linhas extras (manter so 1) "
        "antes da entrega - nao e algo pra corrigir automaticamente sem revisar caso a caso "
        "(podem ter dados diferentes entre as copias, vale conferir)."
    )


if __name__ == "__main__":
    main()
