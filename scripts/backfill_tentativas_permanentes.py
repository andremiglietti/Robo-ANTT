"""
Backfill pontual (24/09/2026, pedido do usuário): o contador `tentativas`
(ver checkpoint.marcar_falha()/falha_provavelmente_permanente(), criado no
mesmo dia) só começa a contar a partir de hoje - falhas salvas ANTES da
correção viraram `tentativas=0` na migração automática, mesmo pros ~181
documentos já confirmados como "problema permanente do servidor" ao longo
de vários dias de evidência real (19-24/09/2026, ver CLAUDE.md).

Sem esse backfill, o filtro que evita gastar 3 rodadas de retentativa por
execução em falha já conhecida só passaria a valer de verdade depois de
mais 2-3 execuções "esquentando" o contador do zero - desperdiçando
exatamente o tempo que a otimização de hoje existe pra evitar.

Este script marca `tentativas = LIMITE_TENTATIVAS_PROVAVEL_PERMANENTE`
direto nos autos que já aparecem na lista exportada por
exportar_falhas_pendentes.py (Falhas_Pendentes_Revisao_Manual.xlsx) - essa
lista já É a evidência real consolidada, não uma suposição nova.

⚠️ CUIDADO DE CONCORRÊNCIA: só é seguro rodar contra um checkpoint cujo
processo (worker) NÃO esteja mais rodando - um worker ativo mantém seu
próprio estado em memória e pode sobrescrever esta edição na próxima vez
que salvar. Por padrão, este script PULA (com aviso) qualquer checkpoint
cuja trava (.lock) ainda esteja com um PID vivo - mesma checagem já usada
em io_seguro.adquirir_trava().

100% local - lê só arquivos já em disco, sem sessão nem contato com o portal.

USO:
    python scripts/backfill_tentativas_permanentes.py           # aplica de verdade
    python scripts/backfill_tentativas_permanentes.py --dry-run  # só mostra o que faria
"""
import argparse
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR / "src"))

from robo_antt import checkpoint as checkpoint_mod  # noqa: E402
from robo_antt.config import OUTPUT_DIR, SHAREPOINT_DIR  # noqa: E402

LISTA_FALHAS_CONHECIDAS = SHAREPOINT_DIR / "Falhas_Pendentes_Revisao_Manual.xlsx"


def _ler_autos_confirmados() -> set[str]:
    wb = load_workbook(str(LISTA_FALHAS_CONHECIDAS), read_only=True)
    ws = wb.active
    return {str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}


def _checkpoint_em_uso(caminho: Path) -> bool:
    """Mesma checagem de io_seguro._processo_ainda_rodando(), sem usar a
    trava de verdade (não queremos ADQUIRIR a trava aqui, só verificar se
    tem dono vivo, e sem interferir num worker que ainda está rodando)."""
    trava = caminho.with_name(caminho.name + ".lock")
    if not trava.exists():
        return False
    try:
        pid = int(trava.read_text(encoding="utf-8").strip())
    except (ValueError, OSError):
        return False  # conteúdo ilegível - trata como "sem dono confirmado", não bloqueia o backfill
    resultado = subprocess.run(["tasklist", "/FI", f"PID eq {pid}"], capture_output=True, text=True)
    return str(pid) in resultado.stdout


def _todos_os_checkpoints() -> list[Path]:
    caminhos = list(OUTPUT_DIR.glob("checkpoint_worker_*.json"))
    if checkpoint_mod.CHECKPOINT_FILE.exists():
        caminhos.append(checkpoint_mod.CHECKPOINT_FILE)
    return caminhos


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dry-run", action="store_true", help="só mostra o que faria, não grava nada")
    args = parser.parse_args()

    if not LISTA_FALHAS_CONHECIDAS.exists():
        print(f"[ATENÇÃO] Lista não encontrada: {LISTA_FALHAS_CONHECIDAS} - rode exportar_falhas_pendentes.py primeiro.")
        return

    autos_confirmados = _ler_autos_confirmados()
    print(f"Autos confirmados como permanentes (lista exportada): {len(autos_confirmados)}")

    total_marcados = 0
    for caminho in sorted(_todos_os_checkpoints()):
        if _checkpoint_em_uso(caminho):
            print(f"  [PULADO] {caminho.name} - trava ativa (worker ainda rodando), inseguro editar agora.")
            continue

        estado = json.loads(caminho.read_text(encoding="utf-8"))
        marcados_aqui = 0
        for auto, info in estado.get("falhas", {}).items():
            if not isinstance(info, dict):
                continue
            if auto not in autos_confirmados:
                continue
            atual = info.get("tentativas", 0)
            if atual < checkpoint_mod.LIMITE_TENTATIVAS_PROVAVEL_PERMANENTE:
                info["tentativas"] = checkpoint_mod.LIMITE_TENTATIVAS_PROVAVEL_PERMANENTE
                marcados_aqui += 1

        if marcados_aqui:
            print(f"  {caminho.name}: {marcados_aqui} auto(s) marcado(s) como permanente confirmado")
            if not args.dry_run:
                caminho_tmp = caminho.with_name(f"_tmp_{caminho.name}")
                caminho_tmp.write_text(json.dumps(estado, ensure_ascii=False, indent=2), encoding="utf-8")
                caminho_tmp.replace(caminho)
        total_marcados += marcados_aqui

    print(f"\nTotal: {total_marcados} falha(s) marcada(s) como permanente confirmado" + (" (dry-run, nada gravado)" if args.dry_run else "."))


if __name__ == "__main__":
    main()
