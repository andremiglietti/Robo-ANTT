"""
Consolida as planilhas-fatia de cada worker na planilha real do SharePoint
(arquitetura de múltiplos workers, 17/09/2026 - ver CLAUDE.md e o plano da
sessão).

Por que existe: cada worker escreve numa planilha LOCAL própria
(data/output/planilha_worker_{N}.xlsx), nunca direto na planilha do
SharePoint - se vários processos escrevessem no mesmo arquivo ao mesmo
tempo, um escreveria por cima do outro (perda de dados). Este script é o
ÚNICO lugar que escreve na planilha real - por isso não tem concorrência.

Reaproveita ja_registrado()/adicionar_registro()/salvar() de planilha.py -
mesma checagem de duplicidade por Auto de Infração usada no resto do robô -
então é seguro rodar quantas vezes quiser (idempotente), inclusive durante
uma execução ainda em andamento dos workers.

USO:
    python scripts/consolidar_planilhas.py
    python scripts/consolidar_planilhas.py --workers 5   (padrão: procura workers 0-9)
"""
import argparse
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR / "src"))

from robo_antt import config  # noqa: E402
from robo_antt.planilha import abrir_ou_criar, adicionar_registro, ja_registrado, registro_da_linha, salvar  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _linhas_da_planilha_worker(caminho: Path) -> list[dict]:
    wb = load_workbook(str(caminho))
    ws = wb.active
    registros = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not any(linha):
            continue
        registros.append(registro_da_linha(linha))
    return registros


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument(
        "--workers", type=int, default=10, help="Quantos IDs de worker procurar (0 até este número -1). Padrão: 10."
    )
    args = parser.parse_args()

    wb_principal = abrir_ou_criar(config.PLANILHA_PATH)
    total_mesclado = 0

    for worker_id in range(args.workers):
        caminho_worker = config.planilha_worker(worker_id)
        if not caminho_worker.exists():
            continue

        registros = _linhas_da_planilha_worker(caminho_worker)
        novos_desse_worker = 0
        for registro in registros:
            auto = registro.get("auto_infracao")
            if not auto:
                continue
            if not ja_registrado(wb_principal, auto):
                adicionar_registro(wb_principal, registro)
                novos_desse_worker += 1

        print(f"Worker {worker_id} ({caminho_worker.name}): {len(registros)} linha(s) na fatia, {novos_desse_worker} nova(s) mesclada(s).")
        total_mesclado += novos_desse_worker

    if total_mesclado:
        salvar(wb_principal, config.PLANILHA_PATH)
        print(f"\n{total_mesclado} linha(s) nova(s) mesclada(s) na planilha final: {config.PLANILHA_PATH}")
    else:
        print("\nNenhuma linha nova pra mesclar (planilha final já estava em dia).")


if __name__ == "__main__":
    main()
