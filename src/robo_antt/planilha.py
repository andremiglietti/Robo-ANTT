"""
Geração/atualização da planilha Excel final com os dados de todas as multas.

Uso típico (ver CLAUDE.md, item 7 da arquitetura):
    wb = abrir_ou_criar(caminho)
    for cada multa nova encontrada:
        if not ja_registrado(wb, registro["auto_infracao"]):
            adicionar_registro(wb, registro)
    salvar(wb, caminho)  # só grava no final, com o cuidado de arquivo temporário
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from robo_antt.io_seguro import substituir_com_retentativa

NOME_ABA = "Multas"

# Ordem e nomes das colunas definidos com o usuário em 15/09/2026 (ver
# CLAUDE.md, item 7 da arquitetura) - não reordenar sem necessidade, já que
# a pessoa que usa a planilha pode ter filtros/fórmulas montados em cima
# dessa ordem depois da entrega.
COLUNAS = [
    "Link do Arquivo",
    "Número do Processo",
    "Auto de Infração",
    "Tipo de Multa",
    "CNPJ",
    "Data Autuação",
    "Data Emissão Doc. Fiscal",
    "Data Emissão Notificação",
    "Data Emissão Boleto",
    "Data Vencimento",
    "Valor",
    "Valor Desconto",
    "Placa",
    "Descrição da Infração",
    "Código de Barras",
    # "Situação" (22/09/2026, pedido do usuário): vem de graça da tabela de
    # busca (row["situacao"] em portal.ler_pagina_atual(), já capturada há
    # dias mas descartada até agora). Acrescentada no FIM da lista (não
    # reordenada pro meio) de propósito - não mexe na posição de nenhuma
    # coluna já existente, então planilhas/filtros que a empresa já tenha
    # montado em cima da ordem antiga continuam funcionando.
    # ✅ Validado ao vivo (mesmo dia): amostra real mostrou 15 valores
    # distintos - todo valor começando com "Arquivado" (Pago/Cancelado) é
    # fechado, qualquer outro (Congelado, Notificação, Auto inscrito na
    # Serasa, Recurso em julgamento, etc.) é aberto/em andamento - dá pra
    # filtrar "multas ainda abertas" só com isso, sem cruzar com outra tela.
    "Situação",
]

# Mapeia o nome da coluna pro nome do campo usado no resto do código
# (portal.py, download.py, extracao.py já usam esses nomes de campo).
_CAMPO_POR_COLUNA = {
    "Link do Arquivo": "link_arquivo",
    "Número do Processo": "numero_processo",
    "Auto de Infração": "auto_infracao",
    "Tipo de Multa": "tipo_multa",
    "CNPJ": "cnpj",
    "Data Autuação": "data_autuacao",
    "Data Emissão Doc. Fiscal": "data_emissao_doc_fiscal",
    "Data Emissão Notificação": "data_emissao_notificacao",
    "Data Emissão Boleto": "data_emissao_boleto",
    "Data Vencimento": "data_vencimento",
    "Valor": "valor",
    "Valor Desconto": "valor_desconto",
    "Placa": "placa",
    "Descrição da Infração": "descricao_infracao",
    "Código de Barras": "codigo_barras",
    "Situação": "situacao",
}

_COLUNA_AUTO_INFRACAO = COLUNAS.index("Auto de Infração") + 1  # openpyxl é 1-based


def abrir_ou_criar(caminho: Path) -> Workbook:
    """Abre a planilha existente (preservando o que já tem), ou cria uma nova
    com o cabeçalho, se ainda não existir."""
    if caminho.exists():
        wb = load_workbook(str(caminho))
        _completar_cabecalho(_aba(wb))
        return wb

    wb = Workbook()
    ws = wb.active
    ws.title = NOME_ABA
    ws.append(COLUNAS)
    return wb


def _completar_cabecalho(ws: Worksheet) -> None:
    """Se a planilha já existia de antes de uma coluna nova ser acrescentada
    em COLUNAS (ex.: 'Situação', 22/09/2026), completa o cabeçalho da aba já
    aberta com as colunas que ainda faltam - sem mexer nas linhas já
    gravadas (ficam em branco nessa coluna nova até serem reprocessadas).
    Não faz nada se o cabeçalho já está em dia."""
    cabecalho_atual = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    proxima_coluna = len(cabecalho_atual) + 1
    for coluna in COLUNAS:
        if coluna not in cabecalho_atual:
            ws.cell(row=1, column=proxima_coluna, value=coluna)
            proxima_coluna += 1


def _aba(wb: Workbook) -> Worksheet:
    if NOME_ABA in wb.sheetnames:
        return wb[NOME_ABA]
    return wb.active


# Nome do atributo usado pra pendurar o cache de autos já conhecidos
# direto no objeto Workbook (ver ja_registrado()/adicionar_registro()) -
# prefixado pra não colidir com nenhum atributo interno do openpyxl.
_ATTR_CACHE_AUTOS = "_robo_antt_cache_autos"


def ja_registrado(wb: Workbook, auto_infracao: str) -> bool:
    """Confere se esse auto já está na planilha, pra não duplicar linha numa
    execução futura do robô (a checagem de arquivo em download.py evita
    baixar o PDF de novo; esta aqui evita duplicar a linha na planilha,
    inclusive se o PDF já existia de uma execução anterior).

    ⚠️ Otimização em 22/09/2026 (ver CLAUDE.md): antes, cada chamada
    reescaneava TODAS as linhas já existentes (O(n) por auto checado) -
    numa consolidação com milhares de linhas dos dois lados
    (scripts/consolidar_planilhas.py), isso vira O(n×m) e fica lento à
    toa. Agora o primeiro auto checado nesta `wb` monta um `set` uma
    única vez (guardado no próprio objeto Workbook, então cada arquivo
    aberto tem seu cache independente) - checagens seguintes na mesma
    `wb` são O(1). `adicionar_registro()` mantém o cache atualizado
    incrementalmente, sem precisar reconstruir do zero."""
    cache = getattr(wb, _ATTR_CACHE_AUTOS, None)
    if cache is None:
        ws = _aba(wb)
        cache = {
            linha[0].value
            for linha in ws.iter_rows(min_row=2, min_col=_COLUNA_AUTO_INFRACAO, max_col=_COLUNA_AUTO_INFRACAO)
            if linha[0].value is not None
        }
        setattr(wb, _ATTR_CACHE_AUTOS, cache)
    return auto_infracao in cache


def registro_da_linha(linha: tuple) -> dict:
    """Inverso de adicionar_registro(): dado os valores de uma linha da
    planilha (na ordem de COLUNAS), devolve o dict com as chaves de campo
    (auto_infracao, cnpj, valor, etc.). Usado por
    scripts/consolidar_planilhas.py pra reler as planilhas-fatia de cada
    worker (arquitetura de múltiplos workers, 17/09/2026) e mesclar na
    planilha final com adicionar_registro()."""
    return {_CAMPO_POR_COLUNA[coluna]: valor for coluna, valor in zip(COLUNAS, linha)}


def adicionar_registro(wb: Workbook, registro: dict) -> None:
    """Adiciona uma linha nova com os dados de uma multa. `registro` é um
    dict com as chaves em _CAMPO_POR_COLUNA.values() (auto_infracao, cnpj,
    valor, etc.) - chaves ausentes viram célula em branco, não erro."""
    ws = _aba(wb)
    linha = [registro.get(_CAMPO_POR_COLUNA[coluna]) for coluna in COLUNAS]
    ws.append(linha)
    # mantém o cache de ja_registrado() em dia (só se ele já existir - se
    # ninguém chamou ja_registrado() nesta wb ainda, não há cache pra
    # atualizar, e tudo bem: ele nasce correto na 1ª chamada futura, já
    # incluindo esta linha).
    cache = getattr(wb, _ATTR_CACHE_AUTOS, None)
    if cache is not None:
        cache.add(registro.get("auto_infracao"))


def salvar(wb: Workbook, caminho_final: Path) -> None:
    """Salva a planilha com o mesmo cuidado da gravação de qualquer arquivo
    na pasta sincronizada com o SharePoint (ver CLAUDE.md, "Cuidado técnico
    na gravação"): escreve num arquivo temporário primeiro e só then troca
    pelo arquivo final, pra nunca deixar a pasta com um .xlsx pela metade
    caso o OneDrive tente sincronizar no meio da escrita.
    """
    caminho_final.parent.mkdir(parents=True, exist_ok=True)
    caminho_tmp = caminho_final.with_name(f"_tmp_{caminho_final.name}")
    wb.save(str(caminho_tmp))
    substituir_com_retentativa(caminho_tmp, caminho_final)
